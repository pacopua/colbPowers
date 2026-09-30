import re
from datetime import date
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

from .config import StyleConfig, MeetingMetadata


# ---------------------------------------------------------------------------
# Color constants for tables
# ---------------------------------------------------------------------------
COLBAI_HEADER_COLOR = "D5E8F0"   # Colbai table header (light blue)
CLIENT_HEADER_COLOR = "D9EAD3"   # Client table header (light green)
STATUS_COLORS = {
    "pendiente": "F2F2F2",
    "en curso":  "FFF2CC",
    "hecho":     "D9EAD3",
}


# ---------------------------------------------------------------------------
# Base helpers (unchanged from original)
# ---------------------------------------------------------------------------

def _hex_to_rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_run(run, font_name: str, size: int, color: str, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.color.rgb = _hex_to_rgb(color)
    run.font.bold = bold


def _add_heading(doc: Document, text: str, level: int, cfg: StyleConfig) -> None:
    para = doc.add_paragraph()
    if level == 0:
        font_name, size, color = cfg.heading_font, cfg.title_size, cfg.title_color
    elif level == 1:
        font_name, size, color = cfg.heading_font, cfg.heading1_size, cfg.heading1_color
    elif level == 2:
        font_name, size, color = cfg.heading_font, cfg.heading2_size, cfg.heading2_color
    else:
        font_name, size, color = cfg.heading_font, cfg.heading3_size, cfg.heading3_color
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            _set_run(run, font_name, size, color, bold=True)
        else:
            run = para.add_run(part)
            _set_run(run, font_name, size, color, bold=True)


def _add_body(doc: Document, text: str, cfg: StyleConfig, bullet: bool = False) -> None:
    para = doc.add_paragraph(style="List Bullet" if bullet else "Normal")
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            _set_run(run, cfg.body_font, cfg.body_size, cfg.body_color, bold=True)
        else:
            run = para.add_run(part)
            _set_run(run, cfg.body_font, cfg.body_size, cfg.body_color)


# ---------------------------------------------------------------------------
# Cover page helpers
# ---------------------------------------------------------------------------

def _centered_para(doc, text, font_name, size, color, bold=False):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    _set_run(run, font_name, size, color, bold)
    return para


def _add_cover_page(doc: Document, metadata: MeetingMetadata, cfg: StyleConfig) -> None:
    created_date = metadata.created_date or date.today().strftime("%d/%m/%Y")

    # Vertical spacing
    for _ in range(4):
        doc.add_paragraph()

    # Client name
    _centered_para(doc, metadata.client_name, cfg.heading_font, 28, "4472C4", bold=True)

    # Subtitle
    _centered_para(doc, "Reunión de seguimiento", cfg.heading_font, 16, "595959", bold=False)

    # Blank paragraph
    doc.add_paragraph()

    # Next meeting info
    _centered_para(
        doc,
        f"Próxima reunión: {metadata.next_meeting_date} | Hora: {metadata.next_meeting_time}",
        cfg.body_font, 11, "404040",
    )

    # Participants (only if list is non-empty)
    if metadata.participants:
        _centered_para(
            doc,
            f"Participantes: {', '.join(metadata.participants)}",
            cfg.body_font, 11, "404040",
        )

    # Created date
    _centered_para(doc, f"Creado: {created_date}", cfg.body_font, 11, "404040")

    # Page break
    para = doc.add_paragraph()
    run = para.add_run()
    run.add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Header / footer helpers
# ---------------------------------------------------------------------------

def _add_bottom_border(para, color="4472C4", size=12):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _set_right_tab(para, position_cm=15.0):
    pos_twips = int(position_cm * 567.17)
    pPr = para._p.get_or_add_pPr()
    tabs_el = OxmlElement('w:tabs')
    tab_el = OxmlElement('w:tab')
    tab_el.set(qn('w:val'), 'right')
    tab_el.set(qn('w:pos'), str(pos_twips))
    tabs_el.append(tab_el)
    pPr.append(tabs_el)


def _add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    run._r.append(fldChar1)
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' PAGE '
    run._r.append(instrText)
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar2)


def _add_header_footer(section, metadata: MeetingMetadata, cfg: StyleConfig) -> None:
    created_date = metadata.created_date or date.today().strftime("%d/%m/%Y")

    # --- Header ---
    header = section.header
    # Clear default paragraphs
    for para in header.paragraphs:
        para.clear()

    # Paragraph 1: Client name
    h_para1 = header.paragraphs[0]
    h_para1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = h_para1.add_run(metadata.client_name)
    _set_run(run1, cfg.heading_font, 11, "4472C4", bold=True)

    # Paragraph 2: subtitle
    h_para2 = header.add_paragraph()
    h_para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = h_para2.add_run(
        f"Reunión de seguimiento – {metadata.next_meeting_date}"
    )
    _set_run(run2, cfg.body_font, 10, "595959")

    # Paragraph 3: empty with bottom border (blue separator)
    h_para3 = header.add_paragraph()
    _add_bottom_border(h_para3)

    # --- Footer ---
    footer = section.footer
    footer_para = footer.paragraphs[0]
    _set_right_tab(footer_para, position_cm=14.5)

    left_run = footer_para.add_run(
        f"Colbai – Confidencial | Creado: {created_date}"
    )
    _set_run(left_run, cfg.body_font, 9, "595959")

    footer_para.add_run("\t")

    pg_run = footer_para.add_run()
    _add_page_number(pg_run)
    _set_run(pg_run, cfg.body_font, 9, "595959")


# ---------------------------------------------------------------------------
# Markdown table parsing
# ---------------------------------------------------------------------------

def _parse_markdown_table(lines, start_idx):
    """
    Returns (rows, next_idx) where rows is list[list[str]].
    Skips separator rows (|---|---|).
    """
    rows = []
    i = start_idx
    while i < len(lines):
        line = lines[i].strip()
        if not line.startswith("|"):
            break
        if re.match(r"^\|[\s\-|]+\|$", line):  # separator row
            i += 1
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


# ---------------------------------------------------------------------------
# Colored table rendering
# ---------------------------------------------------------------------------

def _set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def _add_colored_table(doc: Document, rows: list, is_colbai: bool, cfg: StyleConfig) -> None:
    if not rows:
        return
    header_bg = COLBAI_HEADER_COLOR if is_colbai else CLIENT_HEADER_COLOR
    n_cols = len(rows[0])
    table = doc.add_table(rows=0, cols=n_cols)
    table.style = "Table Grid"

    for row_idx, cells_data in enumerate(rows):
        row = table.add_row()
        is_header = (row_idx == 0)
        for col_idx, cell_text in enumerate(cells_data):
            cell = row.cells[col_idx]
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(cell_text)
            _set_run(run, cfg.body_font, cfg.body_size, cfg.body_color, bold=is_header)
            # Set cell background
            if is_header:
                bg_color = header_bg
            else:
                if col_idx == 1:
                    status_key = cell_text.strip().lower()
                    bg_color = STATUS_COLORS.get(status_key, "FFFFFF")
                else:
                    bg_color = "FFFFFF"
            _set_cell_bg(cell, bg_color)

    doc.add_paragraph()  # spacing after table


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def export_to_docx(
    markdown_text: str,
    style_config: StyleConfig = None,
    metadata: MeetingMetadata = None,
) -> bytes:
    cfg = style_config or StyleConfig()
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(cfg.margin_cm)
        section.bottom_margin = Cm(cfg.margin_cm)
        section.left_margin = Cm(cfg.margin_cm)
        section.right_margin = Cm(cfg.margin_cm)

    # Cover page (only when metadata provided)
    if metadata is not None:
        _add_cover_page(doc, metadata, cfg)

    # Header / footer (only when metadata provided)
    if metadata is not None:
        for section in doc.sections:
            _add_header_footer(section, metadata, cfg)

    # Main content parsing loop
    lines = markdown_text.splitlines()
    i = 0
    last_h2_text = ""
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("#### "):
            _add_heading(doc, stripped[5:], level=3, cfg=cfg)
        elif stripped.startswith("### "):
            _add_heading(doc, stripped[4:], level=2, cfg=cfg)
        elif stripped.startswith("## "):
            last_h2_text = stripped[3:]
            _add_heading(doc, stripped[3:], level=1, cfg=cfg)
        elif stripped.startswith("# "):
            _add_heading(doc, stripped[2:], level=0, cfg=cfg)
        elif stripped.startswith("| ") or stripped.startswith("|"):
            rows, i = _parse_markdown_table(lines, i)
            is_colbai = "colbai" in last_h2_text.lower()
            _add_colored_table(doc, rows, is_colbai, cfg)
            continue  # i already advanced
        elif stripped.startswith("- ") or stripped.startswith("* "):
            _add_body(doc, stripped[2:], cfg, bullet=True)
        else:
            _add_body(doc, stripped, cfg)
        i += 1

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
