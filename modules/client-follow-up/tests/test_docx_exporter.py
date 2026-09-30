from io import BytesIO
import pytest
from docx import Document
from meeting_summarizer.config import StyleConfig, MeetingMetadata
from meeting_summarizer.docx_exporter import export_to_docx


def _load(docx_bytes: bytes) -> Document:
    return Document(BytesIO(docx_bytes))


def _texts(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs if p.text.strip()]


def test_returns_bytes():
    result = export_to_docx("# Title\n\n## Section\n\nSome text.")
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_headings_and_body_present():
    md = "# Meeting Title\n\n## Attendees\n\nAlice and Bob.\n\n### Sub-section\n\nDetail here."
    doc = _load(export_to_docx(md))
    texts = _texts(doc)
    assert "Meeting Title" in texts
    assert "Attendees" in texts
    assert "Sub-section" in texts
    assert "Alice and Bob." in texts


def test_bullet_items_present():
    md = "## Action Items\n\n- Alice: write report\n- Bob: review code"
    doc = _load(export_to_docx(md))
    texts = _texts(doc)
    assert "Alice: write report" in texts
    assert "Bob: review code" in texts


def test_bold_inline_creates_bold_run():
    md = "Some **bold** text here."
    doc = _load(export_to_docx(md))
    body_para = next(p for p in doc.paragraphs if "bold" in p.text)
    bold_run = next(r for r in body_para.runs if r.text == "bold")
    assert bold_run.bold is True


def test_custom_style_config_applied():
    config = StyleConfig(body_font="Times New Roman", title_size=30)
    md = "# Big Title\n\nNormal text."
    doc = _load(export_to_docx(md, config))
    title_para = next(p for p in doc.paragraphs if p.text == "Big Title")
    assert title_para.runs[0].font.size.pt == 30


def test_empty_markdown_returns_valid_docx():
    result = export_to_docx("")
    assert isinstance(result, bytes)
    doc = _load(result)
    assert doc is not None


def test_default_style_config_used_when_none():
    result = export_to_docx("# Title", style_config=None)
    doc = _load(result)
    title_para = next(p for p in doc.paragraphs if p.text == "Title")
    assert title_para.runs[0].font.size.pt == 24  # StyleConfig default


def test_cover_page_contains_client_name():
    meta = MeetingMetadata(client_name="Acme Corp", next_meeting_date="01/06/2026")
    result = export_to_docx("# Title\n\nBody text.", metadata=meta)
    doc = _load(result)
    texts = _texts(doc)
    assert any("Acme Corp" in t for t in texts)


def test_table_in_docx():
    md = "## Estado de Tareas – Colbai\n| Tarea | Estado | Prioridad | Notas |\n|---|---|---|---|\n| Fix bug | Hecho | Alta | |\n"
    result = export_to_docx(md)
    doc = _load(result)
    # Table should exist
    assert len(doc.tables) == 1
    # Header row should have "Tarea"
    assert doc.tables[0].rows[0].cells[0].text == "Tarea"


def test_table_status_cell_color():
    md = "## Estado de Tareas – Colbai\n| Tarea | Estado | Prioridad | Notas |\n|---|---|---|---|\n| Fix bug | Hecho | Alta | |\n"
    result = export_to_docx(md)
    doc = _load(result)
    table = doc.tables[0]
    # Data row's Estado cell (col 1) should have green background (D9EAD3)
    status_cell = table.rows[1].cells[1]
    tc = status_cell._tc
    shd = tc.find('.//{%s}shd' % 'http://schemas.openxmlformats.org/wordprocessingml/2006/main')
    fill = shd.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill')
    assert fill.upper() == "D9EAD3"
